"""Build a simple, clinician-facing psychiatric symptom-category survey workbook."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break


OUTPUT_PATH = Path("docs/anketler/psikiyatrik_semptom_kategori_uzman_anketi_v1_0.xlsx")


# These are descriptive themes, not diagnoses. They match the current project taxonomy.
CATEGORIES = [
    ("Uyku bozukluğu", "Uykuya dalma/sürdürme güçlüğü veya aşırı uyku hali."),
    ("Anksiyete / gerginlik", "Kaygı, endişe, gerilim, panik veya huzursuzluk."),
    ("Depresif belirtiler", "Çökkünlük, keyifsizlik, isteksizlik, enerji azalması veya ağlama."),
    ("Öfke / ajitasyon", "Kolay öfkelenme, öfke patlaması, saldırganlık veya belirgin ajitasyon."),
    ("İstenmeyen / olumsuz düşünceler", "Kişiyi rahatsız eden, istenmeyen veya olumsuz düşünce içeriği."),
    ("Ani duygusal tepkiler", "Ani/uygunsuz gülme-ağlama veya belirgin duygusal dalgalanma."),
    ("Manik / taşkın belirtiler", "Taşkınlık, aşırı enerji, hızlanmış konuşma veya düşünce."),
    ("Halüsinasyon / hezeyan", "Algı yaşantıları veya gerçeklikle uyuşmayan, sabit düşünce içeriği."),
    ("İştah / kilo değişimi", "İştahın artması-azalması ya da belirgin kilo değişimi."),
    ("Sosyal içe kapanma / işlev kaybı", "Sosyal geri çekilme veya günlük işlevlerde belirgin zorlanma."),
    ("Madde kullanımı", "Alkol ya da diğer maddelerin kullanımıyla ilgili ifade."),
    ("Tedaviye uyumsuzluk", "İlaç, randevu veya tedaviyi sürdürme ile ilgili güçlük."),
    ("Somatik yakınmalar", "Mide bulantısı, kusma, çarpıntı, nefes darlığı veya baş ağrısı gibi bedensel yakınmalar."),
]


SYMPTOMS = [
    ("Uyku", "Uykuya dalmakta güçlük"),
    ("Uyku", "Gece boyunca sık uyanma"),
    ("Uyku", "Sabah çok erken uyanma"),
    ("Uyku", "Gün içinde aşırı uyku hali veya uzun süre uyuma"),
    ("Kaygı ve duygudurum", "Sürekli kaygı veya endişe hissetme"),
    ("Kaygı ve duygudurum", "Geleceğe ilişkin yoğun kaygı"),
    ("Kaygı ve duygudurum", "Gerilim veya huzursuzluk hissetme"),
    ("Kaygı ve duygudurum", "Panik hissi ya da ani yoğun korku"),
    ("Kaygı ve duygudurum", "Keyifsiz veya çökkün hissetme"),
    ("Kaygı ve duygudurum", "İsteksizlik veya daha önce zevk aldığı etkinliklerden zevk alamama"),
    ("Kaygı ve duygudurum", "Enerji azalması ya da belirgin yorgunluk"),
    ("Kaygı ve duygudurum", "Ağlama isteği veya sık ağlama"),
    ("Kaygı ve duygudurum", "Kolay sinirlenme"),
    ("Kaygı ve duygudurum", "Öfke patlaması veya saldırgan davranış"),
    ("Kaygı ve duygudurum", "Yerinde duramama veya belirgin ajitasyon"),
    ("Aktivasyon", "Alışılmadık derecede enerjik ya da taşkın hissetme"),
    ("Aktivasyon", "Normalden daha hızlı veya çok konuşma"),
    ("Aktivasyon", "Düşüncelerin hızlandığını hissetme"),
    ("Düşünce ve algı", "İstenmeyen, rahatsız edici düşünceler"),
    ("Düşünce ve algı", "Kendini değersiz ya da yoğun suçlu hissetme"),
    ("Düşünce ve algı", "Yoğun umutsuzluk"),
    ("Düşünce ve algı", "Sesler duyma"),
    ("Düşünce ve algı", "Gerçekte olmayan şeyleri görme"),
    ("Düşünce ve algı", "Başkalarının kendisine zarar vereceğine inanma"),
    ("Düşünce ve algı", "Başkalarının düşüncelerini okuyabildiğini veya düşüncelerini kontrol ettiğini düşünme"),
    ("Düşünce ve algı", "Ani gülme veya ağlama nöbetleri"),
    ("Düşünce ve algı", "Duyguların kısa sürede belirgin biçimde değişmesi"),
    ("Bilişsel", "Dikkat toplamakta zorlanma"),
    ("Bilişsel", "Unutkanlık"),
    ("İştah ve işlev", "İştah azalması"),
    ("İştah ve işlev", "İştah artışı"),
    ("İştah ve işlev", "Belirgin kilo kaybı veya kilo alma"),
    ("İştah ve işlev", "İnsanlardan uzaklaşma veya sosyal ortamlardan kaçınma"),
    ("İştah ve işlev", "Günlük sorumlulukları sürdürmekte zorlanma"),
    ("İştah ve işlev", "İş veya okul performansında belirgin düşüş"),
    ("Bedensel", "Mide bulantısı"),
    ("Bedensel", "Kusma"),
    ("Bedensel", "Çarpıntı"),
    ("Bedensel", "Nefes darlığı"),
    ("Bedensel", "Baş ağrısı"),
    ("Madde ve tedavi", "Alkol veya madde kullanımında artış"),
    ("Madde ve tedavi", "İlaçları düzensiz kullanma ya da bırakma"),
    ("Madde ve tedavi", "Tedaviyi reddetme veya randevulara uymama"),
]


THIN_GRAY = Side(style="thin", color="D9E2F3")
LIGHT_BLUE = PatternFill("solid", fgColor="DCE6F1")
PALE_BLUE = PatternFill("solid", fgColor="EEF4FA")
PALE_YELLOW = PatternFill("solid", fgColor="FFF8E7")
WHITE = PatternFill("solid", fgColor="FFFFFF")


def _bordered(cell, fill=WHITE):
    cell.fill = fill
    cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
    cell.font = Font(name="Aptos", size=10)
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def _section_title(sheet, row, last_letter, title, subtitle):
    sheet.merge_cells(f"A{row}:{last_letter}{row}")
    cell = sheet.cell(row, 1, title)
    cell.font = Font(name="Aptos Display", size=13, bold=True, color="1F4E78")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24
    sheet.merge_cells(f"A{row + 1}:{last_letter}{row + 1}")
    cell = sheet.cell(row + 1, 1, subtitle)
    cell.font = Font(name="Aptos", size=10, color="3F3F3F")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row + 1].height = 30


def build_workbook():
    workbook = Workbook()
    survey = workbook.active
    survey.title = "Anket"

    category_start = 4
    category_end = category_start + len(CATEGORIES) - 1
    category_suggestion_column = category_end + 1
    confidence_column = category_suggestion_column + 1
    note_column = confidence_column + 1
    last_letter = get_column_letter(note_column)

    survey.merge_cells(f"A1:{last_letter}1")
    survey["A1"] = "Psikiyatrik Semptom – Kategori Eşleştirme Uzman Anketi"
    survey["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="1F4E78")
    survey["A1"].alignment = Alignment(vertical="center")
    survey.row_dimensions[1].height = 28

    instructions = [
        "Amaç: Her ifadenin uygun betimleyici klinik tema veya temalarıyla eşleştirilmesi. Bu form tanı koyma veya tedavi önerisi için kullanılmaz.",
        "Nasıl doldurulur: Her satırdaki ifadeye uygun bulduğunuz tüm kategori sütunlarında açılır listeden ✓ seçin. Birden fazla kategori işaretleyebilirsiniz. Hiçbiri uygun değilse kategori hücrelerini boş bırakın.",
        "Uygun kategori yoksa: “Uygun kategori yoksa: kategori önerisi” sütununa kendi kategori adınızı yazın. Listede olmayan bir semptom eklemek için aşağıdaki “Ek semptom” satırlarından birine ifadeyi yazıp aynı şekilde eşleştirin.",
        "Güven puanı: İşaretlediğiniz kategori veya kategorilerin ifadeyle eşleştiğine dair kişisel güveninizdir. 1 = çok düşük, 2 = düşük, 3 = orta, 4 = yüksek, 5 = çok yüksek. Not isteğe bağlıdır; örnek: “İfade tek başına belirsiz, süre ve bağlam bilgisi gerekli.”",
    ]
    for row, text in enumerate(instructions, start=2):
        survey.merge_cells(f"A{row}:{last_letter}{row}")
        cell = survey.cell(row, 1, text)
        cell.font = Font(name="Aptos", size=10, color="3F3F3F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        survey.row_dimensions[row].height = 40 if row in {3, 5} else 32

    headers = ["Grup", "No", "Semptom ifadesi"] + [label for label, _ in CATEGORIES] + ["Uygun kategori yoksa:\nkategori önerisi", "Güven\n(1–5)", "Not\n(isteğe bağlı)"]
    header_row = 6
    for column, label in enumerate(headers, start=1):
        cell = survey.cell(header_row, column, label)
        _bordered(cell, LIGHT_BLUE)
        cell.font = Font(name="Aptos", size=9, bold=True, color="1F4E78")
        if category_start <= column <= category_end:
            cell.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    survey.row_dimensions[header_row].height = 145

    data_start = 7
    for number, (group, phrase) in enumerate(SYMPTOMS, start=1):
        row = data_start + number - 1
        values = [group, number, phrase] + [""] * len(CATEGORIES) + ["", "", ""]
        for column, value in enumerate(values, start=1):
            cell = survey.cell(row, column, value)
            _bordered(cell, PALE_BLUE if number % 2 == 0 else WHITE)
            if category_start <= column <= category_end:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        survey.row_dimensions[row].height = 34
    predefined_end = data_start + len(SYMPTOMS) - 1
    extra_symptom_start = predefined_end + 1
    extra_symptom_end = extra_symptom_start + 9
    for number, row in enumerate(range(extra_symptom_start, extra_symptom_end + 1), start=1):
        values = ["Ek semptom (uzman ekler)", f"Ek {number}", ""] + [""] * len(CATEGORIES) + ["", "", ""]
        for column, value in enumerate(values, start=1):
            cell = survey.cell(row, column, value)
            _bordered(cell, PALE_YELLOW)
            if category_start <= column <= category_end:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        survey.row_dimensions[row].height = 34
    table_end = extra_symptom_end

    tick_validation = DataValidation(type="list", formula1='"✓"', allow_blank=True)
    tick_validation.error = "Bu hücreyi boş bırakın veya açılır listeden ✓ seçin."
    survey.add_data_validation(tick_validation)
    tick_validation.add(f"{get_column_letter(category_start)}{data_start}:{get_column_letter(category_end)}{table_end}")
    confidence_validation = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    survey.add_data_validation(confidence_validation)
    confidence_validation.add(f"{get_column_letter(confidence_column)}{data_start}:{get_column_letter(confidence_column)}{table_end}")

    guide_title_row = table_end + 3
    _section_title(
        survey, guide_title_row, last_letter, "Kategori rehberi", "Tik atarken aşağıdaki kısa açıklamaları kullanın. Kategoriler tanı değil, ifadenin betimlediği temalardır."
    )
    guide_header_row = guide_title_row + 3
    guide_spans = [(1, 5), (6, 13), (14, note_column)]
    for title, (start, end) in zip(["Kategori", "Ne zaman işaretlenir?", "Uzman görüşü (isteğe bağlı)"], guide_spans):
        survey.merge_cells(start_row=guide_header_row, start_column=start, end_row=guide_header_row, end_column=end)
        cell = survey.cell(guide_header_row, start, title)
        _bordered(cell, LIGHT_BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    guide_row = guide_header_row + 1
    for label, description in CATEGORIES:
        survey.merge_cells(start_row=guide_row, start_column=1, end_row=guide_row, end_column=5)
        survey.merge_cells(start_row=guide_row, start_column=6, end_row=guide_row, end_column=13)
        survey.merge_cells(start_row=guide_row, start_column=14, end_row=guide_row, end_column=note_column)
        survey.cell(guide_row, 1, label)
        survey.cell(guide_row, 6, description)
        for start, _ in guide_spans:
            _bordered(survey.cell(guide_row, start), PALE_BLUE if guide_row % 2 == 0 else WHITE)
        survey.row_dimensions[guide_row].height = 34
        guide_row += 1

    # A4 landscape: every category remains on the same horizontal page.
    # The form can continue vertically over as many pages as needed.
    for letter, width in {"A": 12, "B": 4, "C": 28}.items():
        survey.column_dimensions[letter].width = width
    for column in range(category_start, category_end + 1):
        survey.column_dimensions[get_column_letter(column)].width = 4.6
    survey.column_dimensions[get_column_letter(category_suggestion_column)].width = 15
    survey.column_dimensions[get_column_letter(confidence_column)].width = 7
    survey.column_dimensions[get_column_letter(note_column)].width = 20

    survey.freeze_panes = "D7"
    survey.auto_filter.ref = f"A6:{last_letter}{table_end}"
    survey.sheet_view.showGridLines = False
    survey.page_setup.orientation = "landscape"
    survey.page_setup.paperSize = survey.PAPERSIZE_A4
    survey.page_setup.fitToWidth = 1
    survey.page_setup.fitToHeight = 0
    survey.sheet_properties.pageSetUpPr.fitToPage = True
    # Only the table header repeats. The title and the filling instructions
    # remain on the first printed page to preserve usable writing space.
    survey.print_title_rows = "6:6"
    survey.page_margins.left = 0.15
    survey.page_margins.right = 0.15
    survey.page_margins.top = 0.3
    survey.page_margins.bottom = 0.3
    survey.page_margins.header = 0.1
    survey.page_margins.footer = 0.1
    # Start the reference guide at the top of a fresh A4 page rather than
    # leaving its title at the bottom of the preceding survey page.
    survey.row_breaks.append(Break(id=guide_title_row - 1))
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_workbook()
    checked = load_workbook(OUTPUT_PATH)
    assert checked.sheetnames == ["Anket"]
    assert checked["Anket"].max_column == 3 + len(CATEGORIES) + 3
    print(f"Anket oluşturuldu: {OUTPUT_PATH}")
